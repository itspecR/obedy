from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from accounts.names import display_name
from lunches.excel_report import fill_report
from lunches.excel_style import ALARM, BRAND, CENTER, GRID, LANDSCAPE, LEFT, fill, fit_page, font, put

SUMMARY_SHEET = "Сводка"
SUMMARY_COLUMNS = (
    ("Сотрудник", 34),
    ("Логин", 16),
    ("Обедов", 10),
    ("Нарушений", 12),
    ("Превышений", 13),
    ("Без возврата", 13),
    ("Среднее, мин", 13),
    ("Перебор, мин", 13),
)
ATTENDANCE_COLUMNS = (("Без обеда", 12), ("Отсутствий", 12))


def summary_columns(with_attendance):
    return SUMMARY_COLUMNS + ATTENDANCE_COLUMNS if with_attendance else SUMMARY_COLUMNS


def person_row(stats, with_attendance):
    account = stats.account
    attendance = [stats.missing, stats.absent] if with_attendance else []
    return [
        display_name(account),
        account.login,
        stats.count,
        stats.violations,
        stats.overruns,
        stats.unreturned,
        stats.average_minutes,
        stats.overrun_minutes,
        *attendance,
    ]


def write_summary_titles(sheet, with_attendance):
    for column, (title, width) in enumerate(summary_columns(with_attendance), start=1):
        put(sheet, 1, column, title, font=font(bold=True, color="FFFFFF"), fill=fill(BRAND), alignment=CENTER, border=GRID)
        sheet.column_dimensions[get_column_letter(column)].width = width


def write_summary_row(sheet, row, stats, with_attendance):
    for column, value in enumerate(person_row(stats, with_attendance), start=1):
        cell = put(sheet, row, column, value, alignment=LEFT if column <= 2 else CENTER, border=GRID)
        if stats.violations:
            cell.fill = fill(ALARM)


def fill_summary(sheet, people, with_attendance):
    write_summary_titles(sheet, with_attendance)
    for row, stats in enumerate(people, start=2):
        write_summary_row(sheet, row, stats, with_attendance)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    fit_page(sheet, LANDSCAPE)


def workbook_bytes(export):
    book = Workbook()
    fill_report(book.active, export)
    fill_summary(book.create_sheet(SUMMARY_SHEET), export.people, export.attendance is not None)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
