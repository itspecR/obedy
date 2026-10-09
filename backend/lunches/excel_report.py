from dataclasses import dataclass
from datetime import datetime, timedelta

from openpyxl.utils import get_column_letter

from accounts.names import display_name
from lunches.clock import clock_text, local
from lunches.excel_style import (
    ALARM,
    AMBER,
    BRAND,
    CENTER,
    DATE_FORMAT,
    GREEN,
    GRID,
    HEADER,
    LEFT,
    MUTED,
    PERSON_BAR,
    RED,
    TOTAL_BORDER,
    UNDERLINE,
    WEEKEND,
    PORTRAIT,
    fill,
    fit_page,
    font,
    merged,
    put,
    style_range,
)
from lunches.rules import is_workday
from lunches.statistics import SECONDS_IN_MINUTE, Period, counts_as_overrun, excess_seconds
from lunches.status import MEASURED, LunchStatus, duration_of, is_violation, status_of

REPORT_SHEET = "Отчёт"
TITLE = "Отчёт «Обеды за период»"
BRAND_TEXT = "ОБЕДЫ"
COLUMNS = (("Дата", 12), ("День", 7), ("Ушёл", 10), ("Вернулся", 11), ("Длительность, мин", 14), ("Перебор, мин", 12), ("Статус", 24))
LAST_COLUMN = len(COLUMNS)
WEEKDAYS = ("пн", "вт", "ср", "чт", "пт", "сб", "вс")
EMPTY = "—"
DAY_OFF = "Выходной"
EVERYONE = "Все"
CHOSEN = "Выбрано: {count}"
DAY_FORMAT = "%d.%m.%Y"
PERIOD_TEXT = "{first} — {last}"
TOTAL = "ИТОГО"
VIOLATIONS_TEXT = "Нарушений: {count}"
COUNTS_TEXT = "Дней рабочих: {workdays} · выходных: {days_off} · обедов: {lunches}"
AUTHOR_TEXT = "Составил: {name}"
SIGNATURE = "Ответственное лицо"
SIGNATURE_DATE = "«___» ____________ 20___ г."
SIGNATURE_CAPTIONS = ((3, 4, "должность"), (5, 5, "личная подпись"), (6, 7, "расшифровка подписи"))
STATUS_COLORS = {LunchStatus.ON_TIME: GREEN, LunchStatus.OVERRUN: RED, LunchStatus.UNRETURNED: AMBER, LunchStatus.ONGOING: MUTED}
TITLE_ROW_HEIGHT = 22
PERSON_ROW_HEIGHT = 20


@dataclass(frozen=True)
class Export:
    period: Period
    people: list
    lunches: list
    now: datetime
    rules: object
    chosen: int
    author: str


def days_of(period):
    return [period.first + timedelta(days=offset) for offset in range((period.last - period.first).days + 1)]


def minutes(seconds):
    return round(seconds / SECONDS_IN_MINUTE)


def lunch_minutes(lunch, now):
    return minutes(duration_of(lunch, now).total_seconds()) if status_of(lunch, now) in MEASURED else None


def overrun_minutes(lunch, now):
    return minutes(excess_seconds(lunch, now)) if counts_as_overrun(lunch, now) else None


def returned_at(lunch, now):
    if lunch.ended_at is None or status_of(lunch, now) == LunchStatus.UNRETURNED:
        return EMPTY
    return clock_text(local(lunch.ended_at))


def lunch_values(lunch, now):
    status = status_of(lunch, now)
    return [clock_text(local(lunch.started_at)), returned_at(lunch, now), lunch_minutes(lunch, now), overrun_minutes(lunch, now), status.label]


def empty_values(workday):
    return [EMPTY, EMPTY, None, None, EMPTY if workday else DAY_OFF]


def set_columns(sheet):
    for index, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_title(sheet):
    for row in (1, 2):
        style_range(sheet, row, 1, 2, fill=fill(BRAND))
    put(sheet, 1, 1, BRAND_TEXT, font=font(bold=True, size=14, color="FFFFFF"), alignment=CENTER)
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=2)
    put(sheet, 1, 3, TITLE, font=font(bold=True, size=16), alignment=CENTER)
    sheet.merge_cells(start_row=1, start_column=3, end_row=2, end_column=LAST_COLUMN)
    sheet.row_dimensions[1].height = TITLE_ROW_HEIGHT
    sheet.row_dimensions[2].height = TITLE_ROW_HEIGHT


def write_meta(sheet, export):
    period = PERIOD_TEXT.format(first=export.period.first.strftime(DAY_FORMAT), last=export.period.last.strftime(DAY_FORMAT))
    chosen = CHOSEN.format(count=export.chosen) if export.chosen else EVERYONE
    label = {"font": font(bold=True, color=MUTED), "fill": fill(HEADER), "alignment": CENTER, "border": GRID}
    value = {"font": font(size=12), "alignment": CENTER, "border": GRID}
    boxes = ((1, 2, "Дата составления", local(export.now).strftime(DAY_FORMAT)), (3, 5, "Отчётный период", period), (6, 7, "Сотрудники", chosen))
    for first, last, title, content in boxes:
        merged(sheet, 4, first, last, title, **label)
        merged(sheet, 5, first, last, content, **value)
    merged(sheet, 6, 1, LAST_COLUMN, AUTHOR_TEXT.format(name=export.author), font=font(size=10, color=MUTED, italic=True), alignment=LEFT)
    return 8


def write_column_titles(sheet, row):
    for column, (title, _) in enumerate(COLUMNS, start=1):
        put(sheet, row, column, title, font=font(bold=True), fill=fill(HEADER), alignment=CENTER, border=GRID)


def write_day(sheet, row, day, lunch, export):
    workday = is_workday(export.rules, day)
    values = lunch_values(lunch, export.now) if lunch else empty_values(workday)
    put(sheet, row, 1, day, number_format=DATE_FORMAT, alignment=CENTER, border=GRID)
    put(sheet, row, 2, WEEKDAYS[day.weekday()], alignment=CENTER, border=GRID)
    for column, content in enumerate(values, start=3):
        put(sheet, row, column, content, alignment=CENTER, border=GRID)
    if lunch:
        sheet.cell(row=row, column=LAST_COLUMN).font = font(bold=True, color=STATUS_COLORS[status_of(lunch, export.now)])
    shade = ALARM if lunch and is_violation(lunch, export.now) else (None if workday else WEEKEND)
    if shade:
        for column in range(1, LAST_COLUMN + 1):
            sheet.cell(row=row, column=column).fill = fill(shade)


def write_totals(sheet, row, stats, own, now):
    total = sum(value for value in (lunch_minutes(lunch, now) for lunch in own) if value is not None)
    style = {"font": font(bold=True), "alignment": CENTER, "border": TOTAL_BORDER, "fill": fill(HEADER)}
    merged(sheet, row, 1, 4, TOTAL, **{**style, "alignment": LEFT})
    put(sheet, row, 5, total, **style)
    put(sheet, row, 6, stats.overrun_minutes, **style)
    put(sheet, row, 7, VIOLATIONS_TEXT.format(count=stats.violations), **{**style, "font": font(bold=True, color=RED if stats.violations else GREEN)})


def write_counts(sheet, row, stats, days, export):
    workdays = sum(1 for day in days if is_workday(export.rules, day))
    text = COUNTS_TEXT.format(workdays=workdays, days_off=len(days) - workdays, lunches=stats.count)
    merged(sheet, row, 1, LAST_COLUMN, text, font=font(size=10, color=MUTED), alignment=LEFT)


def write_person(sheet, row, stats, own, export):
    account = stats.account
    days = days_of(export.period)
    by_day = {lunch.day: lunch for lunch in own}
    bar = {"font": font(bold=True, size=12), "fill": fill(PERSON_BAR), "alignment": LEFT, "border": GRID}
    merged(sheet, row, 1, 5, display_name(account), **bar)
    merged(sheet, row, 6, LAST_COLUMN, account.login, **{**bar, "font": font(size=11, color=MUTED)})
    sheet.row_dimensions[row].height = PERSON_ROW_HEIGHT
    write_column_titles(sheet, row + 1)
    for offset, day in enumerate(days, start=row + 2):
        write_day(sheet, offset, day, by_day.get(day), export)
    totals_row = row + 2 + len(days)
    write_totals(sheet, totals_row, stats, own, export.now)
    write_counts(sheet, totals_row + 1, stats, days, export)
    return totals_row + 3


def write_signature(sheet, row):
    merged(sheet, row, 1, 2, SIGNATURE, font=font(bold=True), alignment=LEFT)
    for first, last, caption in SIGNATURE_CAPTIONS:
        merged(sheet, row, first, last, "", border=UNDERLINE)
        merged(sheet, row + 1, first, last, caption, font=font(size=9, color=MUTED), alignment=CENTER)
    merged(sheet, row + 3, 1, 3, SIGNATURE_DATE, alignment=LEFT)


def by_person(lunches):
    groups = {}
    for lunch in lunches:
        groups.setdefault(lunch.account_id, []).append(lunch)
    return groups


def alphabetical(people):
    return sorted(people, key=lambda stats: display_name(stats.account).lower())


def fill_report(sheet, export):
    sheet.title = REPORT_SHEET
    set_columns(sheet)
    write_title(sheet)
    row = write_meta(sheet, export)
    own = by_person(export.lunches)
    for stats in alphabetical(export.people):
        row = write_person(sheet, row, stats, own.get(stats.account.pk, []), export)
    write_signature(sheet, row + 1)
    fit_page(sheet, PORTRAIT)
