from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FORMULA_TYPE = "f"
TEXT_TYPE = "s"
DATE_FORMAT = "DD.MM.YYYY"

BRAND = "1F3A5F"
PERSON_BAR = "DCE6F5"
HEADER = "EEF2F7"
WEEKEND = "F2F2F2"
ALARM = "FDE2E1"
GREEN = "1E7B34"
RED = "B42318"
AMBER = "B54708"
MUTED = "667085"
LINE = "B8C2D1"

THIN = Side(style="thin", color=LINE)
MEDIUM = Side(style="medium", color=BRAND)
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
UNDERLINE = Border(bottom=Side(style="thin", color="000000"))
TOTAL_BORDER = Border(left=THIN, right=THIN, top=MEDIUM, bottom=THIN)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

PORTRAIT = "portrait"
LANDSCAPE = "landscape"


def fill(color):
    return PatternFill("solid", fgColor=color)


def font(bold=False, size=11, color="000000", italic=False):
    return Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)


def keep_as_text(cell):
    if cell.data_type == FORMULA_TYPE:
        cell.data_type = TEXT_TYPE


def put(sheet, row, column, value, **style):
    cell = sheet.cell(row=row, column=column, value=value)
    keep_as_text(cell)
    for name, setting in style.items():
        setattr(cell, name, setting)
    return cell


def style_range(sheet, row, first, last, **style):
    for column in range(first, last + 1):
        for name, setting in style.items():
            setattr(sheet.cell(row=row, column=column), name, setting)


def merged(sheet, row, first, last, value, **style):
    style_range(sheet, row, first, last, **{key: setting for key, setting in style.items() if key in ("border", "fill")})
    cell = put(sheet, row, first, value, **style)
    if last > first:
        sheet.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    return cell


def fit_page(sheet, orientation):
    sheet.page_setup.orientation = orientation
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_view.showGridLines = False
