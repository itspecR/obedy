import json
from datetime import date, time, timedelta
from io import BytesIO
from types import SimpleNamespace

import pytest
from django.test import Client
from openpyxl import load_workbook

import lunches.stats_api
from accounts.models import Account, Role
from lunches.clock import moment_of
from lunches.models import Lunch
from tests.factories import DEFAULT_PASSWORD, make_account

pytestmark = pytest.mark.django_db

JSON = "application/json"
NOW = moment_of(date(2026, 10, 8), time(14, 0))
OCTOBER = "?date_from=2026-10-01&date_to=2026-10-31"


@pytest.fixture(autouse=True)
def frozen(monkeypatch):
    monkeypatch.setattr(lunches.stats_api, "timezone", SimpleNamespace(now=lambda: NOW))


def signed_in(login="hr", role=Role.HR):
    make_account(login=login, role=role, must_change_password=False)
    client = Client()
    client.post("/api/auth/login", json.dumps({"login": login, "password": DEFAULT_PASSWORD}), content_type=JSON)
    return client


def person(login, full_name):
    return Account.objects.create(login=login, full_name=full_name)


def lunch(account, day, minutes, **fields):
    started = moment_of(day, time(12, 0))
    ended = started + timedelta(minutes=minutes) if minutes is not None else None
    return Lunch.objects.create(account=account, day=day, started_at=started, ended_at=ended, limit_minutes=45, **fields)


@pytest.fixture
def october():
    ivanov = person("ivanov", "Иванов Иван")
    petrova = person("petrova", "Петрова Анна")
    sidorov = person("sidorov", "Сидоров Сидор")
    lunch(ivanov, date(2026, 10, 1), 30)
    lunch(ivanov, date(2026, 10, 2), 55)
    lunch(ivanov, date(2026, 10, 5), 360, auto_closed=True)
    lunch(petrova, date(2026, 10, 1), 40)
    lunch(petrova, date(2026, 10, 2), 47)
    lunch(sidorov, date(2026, 10, 1), 20)
    lunch(sidorov, date(2026, 9, 30), 90)
    return ivanov, petrova, sidorov


def test_period_overview_and_worst_first(october):
    body = signed_in().get(f"/api/lunch/stats{OCTOBER}").json()

    assert body["overview"] == {"count": 6, "violations": 3, "average_minutes": 38, "on_time_percent": 50, "people": 3}
    assert [(row["login"], row["violations"], row["overruns"], row["unreturned"], row["overrun_minutes"]) for row in body["people"]] == [
        ("ivanov", 2, 1, 1, 10),
        ("petrova", 1, 1, 0, 2),
        ("sidorov", 0, 0, 0, 0),
    ]
    assert (body["people"][0]["count"], body["people"][0]["average_minutes"]) == (3, 42)
    assert "departments" not in body


def test_ongoing_lunch_is_counted_but_not_measured():
    ivanov = person("ivanov", "Иванов Иван")
    lunch(ivanov, date(2026, 10, 8), None)

    body = signed_in().get(f"/api/lunch/stats{OCTOBER}").json()

    assert body["overview"] == {"count": 1, "violations": 1, "average_minutes": None, "on_time_percent": None, "people": 1}
    assert body["people"][0]["overrun_minutes"] == 75


def test_empty_period_has_no_rows():
    body = signed_in().get("/api/lunch/stats?date_from=2025-01-01&date_to=2025-01-31").json()

    assert body["overview"] == {"count": 0, "violations": 0, "average_minutes": None, "on_time_percent": None, "people": 0}
    assert body["people"] == []


@pytest.mark.parametrize(
    ("query", "message"),
    [
        ("?date_from=2026-10-31&date_to=2026-10-01", "Начало периода позже конца"),
        ("?date_from=2025-09-30&date_to=2026-10-01", "Период — не больше года"),
    ],
)
def test_bad_period_is_explained(query, message):
    response = signed_in().get(f"/api/lunch/stats{query}")

    assert (response.status_code, response.json()["detail"]) == (400, message)


def test_whole_year_is_allowed():
    assert signed_in().get("/api/lunch/stats?date_from=2025-10-02&date_to=2026-10-02").status_code == 200


@pytest.mark.parametrize("path", ["/api/lunch/stats", "/api/lunch/stats/export"])
def test_employee_cannot_see_statistics(path):
    response = signed_in("worker", Role.EMPLOYEE).get(f"{path}{OCTOBER}")

    assert (response.status_code, response.json()["detail"]) == (403, "Статистика доступна HR и администратору")


def test_admin_sees_statistics(october):
    assert signed_in("boss", Role.ADMIN).get(f"/api/lunch/stats{OCTOBER}").status_code == 200


def workbook_of(response):
    return load_workbook(BytesIO(response.content))


def report_rows(book):
    return list(book["Отчёт"].iter_rows(values_only=True))


def row_starting(rows, first):
    return next(index for index, row in enumerate(rows) if row[0] == first)


@pytest.mark.usefixtures("october")
def test_export_is_a_report_by_days_and_a_summary():
    hr = signed_in()
    Account.objects.filter(login="hr").update(full_name="Кадрова Ольга")

    response = hr.get(f"/api/lunch/stats/export{OCTOBER}")
    book = workbook_of(response)
    rows = report_rows(book)

    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response["Content-Disposition"] == 'attachment; filename="obedy_2026-10-01_2026-10-31.xlsx"'
    assert book.sheetnames == ["Отчёт", "Сводка"]
    assert rows[0][2] == "Отчёт «Обеды за период»"
    assert rows[4][:7] == ("08.10.2026", None, "01.10.2026 — 31.10.2026", None, None, "Все", None)
    assert rows[5][0] == "Составил: Кадрова Ольга"
    block = row_starting(rows, "Иванов Иван")
    assert rows[block][5] == "ivanov"
    assert rows[block + 1] == ("Дата", "День", "Ушёл", "Вернулся", "Длительность, мин", "Перебор, мин", "Статус")
    days = rows[block + 2 : block + 33]
    assert days[0][0].date() == date(2026, 10, 1)
    assert days[0][1:] == ("чт", "12:00", "12:30", 30, None, "В пределах лимита")
    assert days[1][1:] == ("пт", "12:00", "12:55", 55, 10, "Превышение")
    assert days[2][1:] == ("сб", "—", "—", None, None, "Выходной")
    assert days[4][1:] == ("пн", "12:00", "—", None, None, "Возврат не отмечен")
    assert days[5][1:] == ("вт", "—", "—", None, None, "—")
    assert rows[block + 33][0] == "ИТОГО"
    assert rows[block + 33][4:] == (85, 10, "Нарушений: 2")
    assert rows[block + 34][0] == "Дней рабочих: 22 · выходных: 9 · обедов: 3"
    assert [row[0] for row in rows if row[0] in ("Иванов Иван", "Петрова Анна", "Сидоров Сидор")] == ["Иванов Иван", "Петрова Анна", "Сидоров Сидор"]
    assert rows[-4][0] == "Ответственное лицо"
    assert rows[-1][0] == "«___» ____________ 20___ г."
    assert list(book["Сводка"].values)[1] == ("Иванов Иван", "ivanov", 3, 2, 1, 1, 42, 10)
    assert book["Сводка"].freeze_panes == "A2"


def test_export_names_the_chosen_count(october):
    ivanov, petrova, _ = october

    rows = report_rows(workbook_of(signed_in().get(f"/api/lunch/stats/export{OCTOBER}&person={ivanov.pk}&person={petrova.pk}")))

    assert rows[4][5] == "Выбрано: 2"
    assert [row[0] for row in rows if row[0] in ("Иванов Иван", "Петрова Анна", "Сидоров Сидор")] == ["Иванов Иван", "Петрова Анна"]


def test_export_keeps_formula_like_text_as_text():
    ivanov = person("ivanov", "=HYPERLINK(\"http://evil\")")
    lunch(ivanov, date(2026, 10, 1), 30)

    book = workbook_of(signed_in().get(f"/api/lunch/stats/export{OCTOBER}"))
    cells = [cell for sheet in book for row in sheet.iter_rows() for cell in row]

    assert all(cell.data_type != "f" for cell in cells)
    assert "=HYPERLINK(\"http://evil\")" in [cell.value for cell in cells]


def test_selected_people_narrow_overview_and_show_their_lunches(october):
    ivanov, petrova, _ = october

    body = signed_in().get(f"/api/lunch/stats{OCTOBER}&person={ivanov.pk}&person={petrova.pk}").json()

    assert body["overview"] == {"count": 5, "violations": 3, "average_minutes": 43, "on_time_percent": 40, "people": 2}
    assert len(body["people"]) == 3
    assert [(row["name"], row["lunch"]["day"]) for row in body["lunches"]][:3] == [
        ("Иванов Иван", "2026-10-01"),
        ("Петрова Анна", "2026-10-01"),
        ("Иванов Иван", "2026-10-02"),
    ]
    assert {row["person_id"] for row in body["lunches"]} == {ivanov.pk, petrova.pk}


def test_without_selection_daily_lunches_are_not_sent(october):
    assert signed_in().get(f"/api/lunch/stats{OCTOBER}").json()["lunches"] == []


def test_export_contains_only_selected_people(october):
    petrova = october[1]

    book = workbook_of(signed_in().get(f"/api/lunch/stats/export{OCTOBER}&person={petrova.pk}"))

    assert [row[1] for row in book["Сводка"].iter_rows(min_row=2, values_only=True)] == ["petrova"]
    assert "Иванов Иван" not in [row[0] for row in report_rows(book)]
